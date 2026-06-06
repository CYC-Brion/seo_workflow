from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FILE = ROOT / "keyword_library" / "seo_keyword_library_full.xlsx"


def row(keyword, category, lane, priority, cluster, targets, notes=""):
    return {
        "keywords": keyword,
        "category": category,
        "lane": lane,
        "priority": priority,
        "cluster": cluster,
        "internal_link_targets": targets,
        "notes": notes,
    }


def plan_row(week, lane, keyword, category, direction, title_angle, targets):
    return {
        "week": week,
        "lane": lane,
        "keyword": keyword,
        "category": category,
        "direction": direction,
        "title_angle": title_angle,
        "internal_link_targets": targets,
    }


def add_keywords(rows, items, category, lane, priority, cluster, targets, notes=""):
    for keyword in items:
        rows.append(row(keyword, category, lane, priority, cluster, targets, notes))


def add_plan_rows(rows, keywords, category, lane, direction, targets, week_start, title_angle=None):
    for offset, keyword in enumerate(keywords):
        angle = title_angle or direction
        rows.append(plan_row(week_start + offset, lane, keyword, category, direction, angle, targets))


def cluster_direction(cluster):
    return {
        "core": "factory overview and conversion intent",
        "floor": "product selection and fit for floor lamp projects",
        "wall": "product selection and fit for wall lamp projects",
        "table": "product selection and fit for table lamp projects",
        "components": "structure, specs, and component applications",
        "process": "process explanation and quality control",
        "procurement": "supplier comparison and buying guide",
        "application": "use cases, project fit, and buyer scenarios",
        "trust": "proof, capability, and sourcing confidence",
    }.get(cluster, "buyer-oriented technical explainer")


def cluster_targets(cluster):
    return {
        "core": "Home; About Us; Services; Contact",
        "floor": "Floor Lamp Sets; Services; Contact",
        "wall": "Wall Lamp Sets; Services; Contact",
        "table": "Table Lamp Sets; Services; Contact",
        "components": "Tubes; Panels; Universal Joints; Machined Parts",
        "process": "Services; About Us; Contact",
        "procurement": "Services; RFQ; About Us; Contact",
        "application": "Home; Services; Contact",
        "trust": "About Us; Services; Contact",
    }.get(cluster, "Home; Services; Contact")


def pillar_variants():
    return {
        "lighting hardware manufacturer": [
            "manufacturing capability overview",
            "factory strength and customization scope",
            "OEM/ODM decision guide",
            "supplier comparison and sourcing fit",
        ],
        "lighting hardware supplier comparison": [
            "supplier audit checklist",
            "how to compare factories on quality and lead time",
            "pricing, MOQ, and collaboration model",
        ],
        "lighting hardware sourcing guide": [
            "buyer checklist for first-time sourcing",
            "how to reduce sourcing risk",
            "sample, lead time, and order process",
        ],
        "OEM vs ODM lighting hardware": [
            "decision guide for product development",
            "cost, lead time, and customization comparison",
            "which model fits different buyers",
        ],
        "custom lighting hardware": [
            "customization workflow from idea to sample",
            "design support and engineering collaboration",
            "prototype, approval, and production transition",
        ],
        "lighting manufacturing process": [
            "step-by-step manufacturing flow",
            "process control and quality checkpoints",
        ],
        "lighting parts quality control": [
            "incoming inspection, in-process checks, and final QC",
            "inspection checklist for buyers",
        ],
        "surface treatment for lighting hardware": [
            "corrosion resistance and finish selection",
            "plating, coating, and wear resistance",
        ],
        "floor lamp parts supplier": [
            "supplier comparison and selection guide",
            "replacement parts and project sourcing",
            "quality checklist for floor lamp projects",
        ],
        "wall lamp parts supplier": [
            "supplier comparison and selection guide",
            "project fit and customization options",
            "quality checklist for wall lamp projects",
        ],
        "table lamp parts supplier": [
            "supplier comparison and selection guide",
            "replacement parts and spare parts planning",
            "quality checklist for table lamp projects",
        ],
    }


def build_workbook():
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    keyword_rows = []

    add_keywords(keyword_rows,
        [
            "lighting hardware manufacturer",
            "lighting hardware factory",
            "lighting accessories manufacturer",
            "lighting accessories supplier",
            "lighting parts manufacturer",
            "lighting parts supplier",
            "lighting hardware supplier",
            "lighting hardware sourcing guide",
            "lighting hardware supplier comparison",
            "lighting sourcing checklist",
            "OEM vs ODM lighting hardware",
            "OEM lighting accessories",
            "ODM lighting accessories",
            "custom lighting hardware",
            "custom lighting parts",
            "how to choose a reliable lamp parts manufacturer",
            "lighting supplier evaluation",
            "supplier audit checklist",
            "factory capability review",
            "factory strength overview",
        ],
        "Technical Guides",
        "pillar",
        "A",
        "core",
        cluster_targets("core"),
    )

    add_keywords(keyword_rows,
        [
            "floor lamp sets",
            "floor lamp parts",
            "floor lamp parts manufacturer",
            "floor lamp parts supplier",
            "floor lamp accessory supplier",
            "floor lamp kit accessory manufacturer",
            "floor lamp bracket hardware",
            "floor lamp support structure",
            "floor lamp mounting hardware",
            "floor lamp replacement parts",
            "floor lamp spare parts",
            "floor lamp parts diagram",
            "floor lamp parts catalog",
            "floor lamp parts review guide",
            "floor lamp installation guide",
            "floor lamp maintenance guide",
        ],
        "Technical Guides",
        "pillar",
        "A",
        "floor",
        cluster_targets("floor"),
    )

    add_keywords(keyword_rows,
        [
            "wall lamp sets",
            "wall lamp parts",
            "wall lamp parts manufacturer",
            "wall lamp parts supplier",
            "wall lamp accessory supplier",
            "wall lamp bracket hardware",
            "wall lamp mounting hardware",
            "wall light mounting hardware",
            "wall lamp mounting bracket",
            "wall lamp replacement parts",
            "wall lamp spare parts",
            "wall lamp parts diagram",
            "wall lamp parts catalog",
            "wall lamp parts review guide",
            "wall lamp installation guide",
            "wall lamp maintenance guide",
        ],
        "Technical Guides",
        "pillar",
        "A",
        "wall",
        cluster_targets("wall"),
    )

    add_keywords(keyword_rows,
        [
            "table lamp sets",
            "table lamp parts",
            "table lamp parts manufacturer",
            "table lamp parts supplier",
            "table lamp accessories parts",
            "table lamp support structure",
            "table lamp replacement parts",
            "table lamp spare parts",
            "table lamp parts kit",
            "table lamp parts diagram",
            "table lamp parts catalog",
            "table lamp parts review guide",
            "table lamp installation guide",
            "table lamp maintenance guide",
            "table lamp base parts",
        ],
        "Technical Guides",
        "pillar",
        "A",
        "table",
        cluster_targets("table"),
    )

    add_keywords(keyword_rows,
        [
            "custom made parts",
            "custom metal hardware parts",
            "custom metal parts fabrication",
            "custom sheet metal parts",
            "custom cut steel parts",
            "machined parts",
            "machined metal parts factory",
            "precision parts",
            "CNC parts",
            "metal parts manufacturing",
            "metal processing machinery parts",
            "aluminum lamp tube components",
            "steel lamp tube supplier",
            "aluminum tube dimensions chart",
            "structural aluminum tube sizes",
            "aluminum tube section properties",
            "tubes",
            "panels",
            "decorative panel",
            "mounting panel",
            "universal joint",
            "universal joint connector",
            "rotating joint",
            "swivel joint",
        ],
        "Technical Guides",
        "ordinary",
        "B",
        "components",
        cluster_targets("components"),
    )

    add_keywords(keyword_rows,
        [
            "lighting manufacturing process",
            "lighting fixture assembly process",
            "lamp manufacturing process",
            "lighting production process",
            "surface treatment for lighting hardware",
            "surface treatment methods",
            "coating adhesion",
            "corrosion resistance",
            "assembly precision",
            "tolerance control",
            "quality control",
            "lighting parts quality control",
            "factory audit checklist",
            "factory audit system",
            "factory production checklist",
            "product testing",
            "certification",
            "compliance",
            "quality inspection checklist",
            "material selection",
            "finish comparison",
            "load-bearing design",
            "incoming inspection",
            "in-process inspection",
            "final inspection",
        ],
        "Technical Guides",
        "ordinary",
        "B",
        "process",
        cluster_targets("process"),
    )

    add_keywords(keyword_rows,
        [
            "lighting hardware sourcing guide",
            "lighting hardware supplier comparison",
            "lighting supplier evaluation",
            "supplier audit checklist",
            "buyer checklist",
            "RFQ guide",
            "MOQ guide",
            "lead time guide",
            "pricing guide",
            "sample request",
            "customization process",
            "order process",
            "collaboration model",
            "project procurement",
            "procurement guide",
            "cost optimization",
            "sourcing risk control",
            "export logistics",
            "packaging and shipping",
            "delivery planning",
            "contract manufacturing",
            "lighting project sourcing",
        ],
        "Industry Insights",
        "ordinary",
        "B",
        "procurement",
        cluster_targets("procurement"),
    )

    add_keywords(keyword_rows,
        [
            "lighting hardware for hotel projects",
            "hospitality lighting hardware",
            "custom hotel parts",
            "contract lighting project",
            "commercial lighting hardware",
            "residential lighting hardware",
            "interior lighting project",
            "lighting accessories for hotels",
            "project lighting hardware",
            "lighting parts case study",
            "factory case study",
            "installation guide",
            "replacement parts",
            "repair parts",
            "maintenance guide",
            "spare parts catalog",
            "after-sales support",
            "project example",
        ],
        "Industry Insights",
        "ordinary",
        "B",
        "application",
        cluster_targets("application"),
    )

    add_keywords(keyword_rows,
        [
            "lighting hardware compliance",
            "lighting hardware testing",
            "factory capability",
            "production capacity",
            "about us",
            "quality assurance",
            "quality standards",
            "patents and R&D",
            "certification documents",
            "inspection report",
            "buyer confidence checklist",
            "factory visit checklist",
            "sample room setup",
            "design support",
            "engineering collaboration",
        ],
        "Industry Insights",
        "ordinary",
        "B",
        "trust",
        cluster_targets("trust"),
    )

    # Deduplicate while preserving order.
    unique_rows = []
    seen = set()
    for item in keyword_rows:
        key = item["keywords"].lower()
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(item)

    plan_rows = []
    week = 1
    for item in unique_rows:
        direction = cluster_direction(item["cluster"])
        plan_rows.append(plan_row(week, item["lane"], item["keywords"], item["category"], direction, direction, item["internal_link_targets"]))
        week += 1

    # Add pillar expansions for the most important keywords.
    variants = pillar_variants()
    pillar_targets = {
        "lighting hardware manufacturer": "Home; About Us; Services; Contact",
        "lighting hardware supplier comparison": "Home; Services; RFQ; Contact",
        "lighting hardware sourcing guide": "Services; RFQ; Contact",
        "OEM vs ODM lighting hardware": "Services; About Us; Contact",
        "custom lighting hardware": "Services; Contact; About Us",
        "lighting manufacturing process": "Services; About Us; Contact",
        "lighting parts quality control": "Services; About Us; Contact",
        "surface treatment for lighting hardware": "Services; About Us; Contact",
        "floor lamp parts supplier": "Floor Lamp Sets; Services; Contact",
        "wall lamp parts supplier": "Wall Lamp Sets; Services; Contact",
        "table lamp parts supplier": "Table Lamp Sets; Services; Contact",
    }

    for keyword, directions in variants.items():
        category = next((item["category"] for item in unique_rows if item["keywords"] == keyword), "Industry Insights")
        targets = pillar_targets.get(keyword, "Home; Services; Contact")
        lane = "pillar"
        for idx, direction in enumerate(directions, start=1):
            plan_rows.append(plan_row(week, lane, keyword, category, direction, direction, targets))
            week += 1

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # Keywords sheet
    ws = wb.create_sheet("keywords")
    keyword_headers = ["keywords", "category", "lane", "priority", "cluster", "internal_link_targets", "notes"]
    ws.append(keyword_headers)
    for item in unique_rows:
        ws.append([item[h] for h in keyword_headers])

    # Plan sheet
    ws = wb.create_sheet("plan")
    plan_headers = ["week", "lane", "keyword", "category", "direction", "title_angle", "internal_link_targets"]
    ws.append(plan_headers)
    for item in plan_rows:
        ws.append([item[h] for h in plan_headers])

    # Title templates sheet
    ws = wb.create_sheet("title_templates")
    ws.append(["template", "angle_type"])
    title_templates = [
        ("What Is {keyword}? A Practical Guide for Buyers and Engineers", "Educational"),
        ("How {keyword} Works: Materials, Structure, and Manufacturing Basics", "Technical"),
        ("{keyword}: 7 Things to Check Before You Order", "Buyer Checklist"),
        ("How to Choose the Right {keyword} for Your Project", "Selection Guide"),
        ("{keyword} Explained: Specs, Quality, and Common Mistakes", "Explainer"),
        ("{keyword} Comparison: What Buyers Should Evaluate First", "Comparison"),
        ("{keyword} Supplier Guide: Quality, Lead Time, and Customization", "Supplier Guide"),
        ("{keyword} for Hotel and Project Use: What to Know", "Application"),
        ("{keyword} QC Checklist for Reliable Procurement", "QC"),
        ("{keyword} Material and Finish Guide for Better Performance", "Materials"),
        ("OEM vs ODM {keyword}: Which Model Fits Your Needs?", "Decision Guide"),
        ("{keyword} Factory Review: Capacity, Process, and Support", "Factory Review"),
        ("How to Reduce Risk When Buying {keyword}", "Risk Control"),
        ("{keyword} Replacement and Maintenance Guide", "After-Sales"),
        ("{keyword} Installation Tips and Common Pitfalls", "Installation"),
        ("{keyword} Case Study: How Buyers Use It in Real Projects", "Case Study"),
    ]
    for template, angle in title_templates:
        ws.append([template, angle])

    # Article angles sheet
    ws = wb.create_sheet("article_angles")
    ws.append(["angle", "recommended_category", "use_case"])
    article_angles = [
        ("Material selection and durability", "Technical Guides", "Compare metals, finishes, and load-bearing needs."),
        ("Structure and load-bearing design", "Technical Guides", "Explain geometry, stability, and mechanical performance."),
        ("Surface treatment and corrosion resistance", "Technical Guides", "Focus on coating, plating, and wear resistance."),
        ("Assembly precision and tolerance control", "Technical Guides", "Show how fit, alignment, and tolerance affect quality."),
        ("QC and inspection process", "Technical Guides", "Cover incoming inspection, in-process checks, and final QC."),
        ("Factory capability and production capacity", "Industry Insights", "Position the supplier as a manufacturing partner."),
        ("Supplier comparison and sourcing fit", "Industry Insights", "Help buyers compare factories and shortlist vendors."),
        ("OEM vs ODM decision guide", "Industry Insights", "Clarify which cooperation model fits the buyer."),
        ("Lead time, MOQ, and pricing", "Industry Insights", "Support procurement decisions with practical constraints."),
        ("Customization workflow", "Industry Insights", "Show how ideas move from sample to mass production."),
        ("Project application and use cases", "Industry Insights", "Link products to hotel, commercial, and contract projects."),
        ("Replacement, installation, and after-sales", "Technical Guides", "Cover maintenance, service, and support content."),
    ]
    for angle, category, use_case in article_angles:
        ws.append([angle, category, use_case])

    # README sheet
    ws = wb.create_sheet("README")
    ws.append(["item", "detail"])
    readme_rows = [
        ("Purpose", "Full keyword library for the lighting hardware factory site."),
        ("How to use", "Import the keywords sheet into the workflow. The plan sheet is optional but recommended for mixed execution."),
        ("Keywords count", str(len(unique_rows))),
        ("Plan rows", str(len(plan_rows))),
        ("Main clusters", "Core, floor, wall, table, components, process, procurement, application, trust."),
        ("Workflow note", "The workflow reads keywords plus metadata. Plan drives pillar priority and direction, while ordinary keywords can still flow through after filtering."),
        ("Internal links", "Each row includes a default internal-link target set for easier article drafting."),
    ]
    for item, detail in readme_rows:
        ws.append([item, detail])

    # Basic formatting
    header_fill = PatternFill("solid", fgColor="E36D2D")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in ws.columns:
            max_len = 0
            column = column_cells[0].column_letter
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_len:
                    max_len = len(value)
            ws.column_dimensions[column].width = min(max_len + 4, 54)

    wb.save(OUTPUT_FILE)
    print(f"Saved {OUTPUT_FILE}")


if __name__ == "__main__":
    build_workbook()
