from __future__ import annotations

from math import ceil
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "keyword_library" / "seo_phase3_long_cycle.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "keyword_library"
DEFAULT_PARTS = 5


def chunk_rows(rows: Sequence[Tuple], parts: int) -> List[List[Tuple]]:
    if parts <= 0:
        raise ValueError("parts must be greater than 0")
    if not rows:
        return []

    chunk_size = ceil(len(rows) / parts)
    return [list(rows[i : i + chunk_size]) for i in range(0, len(rows), chunk_size)]


def copy_rows(source_rows: Iterable[Tuple], worksheet) -> None:
    for row in source_rows:
        worksheet.append(list(row))


def build_part_workbook(
    source_path: Path,
    output_path: Path,
    keyword_rows: Sequence[Tuple],
    plan_rows: Sequence[Tuple],
) -> None:
    source_wb = load_workbook(source_path, read_only=True, data_only=False)

    keywords_ws = source_wb["keywords"]
    plan_ws = source_wb["plan"]
    title_templates_ws = source_wb["title_templates"]
    article_angles_ws = source_wb["article_angles"]
    readme_ws = source_wb["README"]

    target_wb = Workbook()
    default_ws = target_wb.active
    target_wb.remove(default_ws)

    target_keywords_ws = target_wb.create_sheet("keywords")
    target_plan_ws = target_wb.create_sheet("plan")
    target_title_templates_ws = target_wb.create_sheet("title_templates")
    target_article_angles_ws = target_wb.create_sheet("article_angles")
    target_readme_ws = target_wb.create_sheet("README")

    copy_rows(
        (row for row in keywords_ws.iter_rows(values_only=True, min_row=1, max_row=1)),
        target_keywords_ws,
    )
    copy_rows(keyword_rows, target_keywords_ws)

    copy_rows(
        (row for row in plan_ws.iter_rows(values_only=True, min_row=1, max_row=1)),
        target_plan_ws,
    )
    copy_rows(plan_rows, target_plan_ws)

    copy_rows(title_templates_ws.iter_rows(values_only=True), target_title_templates_ws)
    copy_rows(article_angles_ws.iter_rows(values_only=True), target_article_angles_ws)
    copy_rows(readme_ws.iter_rows(values_only=True), target_readme_ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    target_wb.save(output_path)
    source_wb.close()


def split_workbook(
    source_path: Path = DEFAULT_INPUT,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    parts: int = DEFAULT_PARTS,
) -> List[Tuple[Path, int, int]]:
    wb = load_workbook(source_path, read_only=True, data_only=True)

    keyword_rows = list(wb["keywords"].iter_rows(values_only=True, min_row=2))
    plan_rows = list(wb["plan"].iter_rows(values_only=True, min_row=2))
    keyword_lookup = {
        str(row[0]).strip(): row
        for row in keyword_rows
        if row and row[0] is not None and str(row[0]).strip()
    }
    keyword_order = [
        str(row[0]).strip()
        for row in keyword_rows
        if row and row[0] is not None and str(row[0]).strip()
    ]

    plan_chunks = chunk_rows(plan_rows, parts)
    results: List[Tuple[Path, int, int]] = []

    for index, plan_chunk in enumerate(plan_chunks, start=1):
        keyword_set = {
            str(row[2]).strip()
            for row in plan_chunk
            if row and len(row) > 2 and row[2] is not None and str(row[2]).strip()
        }
        keyword_chunk = [
            keyword_lookup[keyword]
            for keyword in keyword_order
            if keyword in keyword_set and keyword in keyword_lookup
        ]

        output_path = output_dir / f"seo_phase3_long_cycle_part{index}.xlsx"
        build_part_workbook(
            source_path=source_path,
            output_path=output_path,
            keyword_rows=keyword_chunk,
            plan_rows=plan_chunk,
        )
        results.append((output_path, len(keyword_chunk), len(plan_chunk)))

    wb.close()
    return results


def main() -> None:
    results = split_workbook()
    print("Created split workbooks:")
    for output_path, keyword_count, plan_count in results:
        print(
            f"- {output_path.name}: keywords={keyword_count}, plan_rows={plan_count}"
        )


if __name__ == "__main__":
    main()
