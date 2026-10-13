"""Generate the remaining keyword plans and schedule each WordPress post hourly.

The script is deliberately resumable: a post is recorded immediately after WordPress
accepts it, so rerunning it only retries items that do not yet have a post ID.
"""

from __future__ import annotations

import json
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from src.enhanced_features.image_integration.editorial_image_library import (  # noqa: E402
    EditorialImageLibrary,
)
from src.modules.content_generator.ai_content_generator import AIContentGenerator  # noqa: E402
from src.modules.title_generator.ai_title_generator import AITitleGenerator  # noqa: E402
from src.modules.wordpress_integration.wordpress_publisher import WordPressPublisher  # noqa: E402


MASTER_WORKBOOK = PROJECT_ROOT / "keyword_library" / "seo_phase3_long_cycle.xlsx"
RESULTS_DIR = PROJECT_ROOT / "output" / "pending_44_hourly_schedule"
RESULTS_FILE = RESULTS_DIR / "schedule_results.json"
LOG_FILE = RESULTS_DIR / "schedule.log"

PILLAR_DIRECTIONS = {
    "core": "supplier evaluation and factory capability verification",
    "floor": "specification, quality checks, and sourcing fit for floor lamp hardware",
    "wall": "specification, quality checks, and sourcing fit for wall lamp hardware",
    "table": "specification, quality checks, and sourcing fit for table lamp hardware",
}
ORDINARY_DIRECTIONS = {
    "components": "technical specifications, component fit, and production verification",
    "procurement": "B2B sourcing, delivery planning, and commercial risk control",
    "application": "application fit, installation planning, and project requirements",
    "trust": "factory capability, company background, and buyer verification",
}


def log(message: str) -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    line = f"[{datetime.now().astimezone().isoformat(timespec='seconds')}] {message}"
    print(line, flush=True)
    with LOG_FILE.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def load_previous_results() -> dict[str, Any]:
    if not RESULTS_FILE.exists():
        return {"started_at": datetime.now(timezone.utc).isoformat(), "articles": []}
    return json.loads(RESULTS_FILE.read_text(encoding="utf-8"))


def save_results(results: dict[str, Any]) -> None:
    RESULTS_FILE.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")


def ai_result_with_retry(node: Any, payload: Any, previous_calls: int, label: str) -> Any:
    """Respect the provider cooldown instead of rapidly falling back to templates."""
    last_error = "AI call did not complete"
    for attempt in range(1, 4):
        result = node.execute(payload)
        if result.status.value == "completed" and node.stats["ai_calls"] > previous_calls:
            return result
        last_error = result.error_message or last_error
        if attempt < 3:
            wait_seconds = 90 * attempt
            log(f"{label} is rate-limited; retrying in {wait_seconds}s.")
            time.sleep(wait_seconds)
    raise RuntimeError(f"{label} did not complete after retries: {last_error}")


def historical_keywords() -> set[str]:
    """Read keywords that were actually accepted by older workflow filter runs."""
    keywords: set[str] = set()
    for workflow_file in (PROJECT_ROOT / "output").glob("workflow_*.json"):
        try:
            data = json.loads(workflow_file.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        steps = data.get("steps", {})
        values = steps.get("3_filtered_keywords", {}).get("data", [])
        titles = steps.get("4_titles", {}).get("data", [])
        keywords.update(str(value).strip().lower() for value in values if str(value).strip())
        # Some legacy runs kept the original topic only in the generated title.
        # Treat those titles as historical coverage to avoid scheduling near-duplicates.
        keywords.update(str(title).strip().lower() for title in titles if str(title).strip())
    return keywords


def remaining_items() -> list[dict[str, str]]:
    workbook = load_workbook(MASTER_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["keywords"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().lower() for value in next(rows)]
    completed_text = "\n".join(historical_keywords())
    items: list[dict[str, str]] = []

    for values in rows:
        row = dict(zip(headers, values))
        keyword = str(row.get("keywords") or "").strip()
        if not keyword or keyword.lower() in completed_text:
            continue
        cluster = str(row.get("cluster") or "").strip().lower()
        lane = str(row.get("lane") or "ordinary").strip().lower()
        direction_map = PILLAR_DIRECTIONS if lane == "pillar" else ORDINARY_DIRECTIONS
        items.append({
            "keyword": keyword,
            "category": str(row.get("category") or "Industry Insights").strip(),
            "lane": lane,
            "priority": str(row.get("priority") or "B").strip(),
            "cluster": cluster,
            "direction": direction_map.get(cluster, "B2B sourcing, specifications, and manufacturing verification"),
            "internal_link_targets": str(row.get("internal_link_targets") or "").strip(),
        })
    return items


def main() -> int:
    results = load_previous_results()
    all_items = remaining_items()
    existing = {entry.get("keyword", "").lower() for entry in results["articles"] if entry.get("post_id")}
    pending_items = [item for item in all_items if item["keyword"].lower() not in existing]
    if not pending_items:
        log("No unscheduled pending keyword plans remain.")
        return 0

    # Keep the hourly grid stable even if generation itself takes more than one hour
    # or the process has to be restarted after a provider cooldown.
    first_slot_text = results.get("first_schedule_utc")
    if first_slot_text:
        first_slot = datetime.fromisoformat(first_slot_text)
    else:
        order = {item["keyword"].lower(): index for index, item in enumerate(all_items)}
        recorded_slots = [
            entry for entry in results["articles"]
            if entry.get("post_id") and entry.get("schedule_utc") and entry.get("keyword", "").lower() in order
        ]
        if recorded_slots:
            anchor = min(recorded_slots, key=lambda entry: entry["schedule_utc"])
            first_slot = datetime.fromisoformat(anchor["schedule_utc"]) - timedelta(
                hours=order[anchor["keyword"].lower()]
            )
        else:
            first_slot = datetime.now(timezone.utc) + timedelta(hours=1)
        results["first_schedule_utc"] = first_slot.isoformat()
        save_results(results)
    library = EditorialImageLibrary(PROJECT_ROOT / "pic" / "editorial_library")
    title_generator = AITitleGenerator("pending_schedule_titles", config={"batch_size": 1, "enable_fallback": False})
    content_generator = AIContentGenerator("pending_schedule_content", config={"batch_size": 1, "target_word_count": 2000, "enable_fallback": False})
    publisher = WordPressPublisher("pending_schedule_publisher", config={
        "batch_size": 1,
        "api_type": "rest",
        "default_status": "draft",
        "auto_create_categories": False,
        "auto_create_tags": False,
        "enable_media_upload": True,
    })

    for index, item in enumerate(all_items):
        if item["keyword"].lower() in existing:
            continue
        slot = first_slot + timedelta(hours=index)
        log(f"Generating slot {index + 1}/{len(all_items)}: {item['keyword']} for {slot.isoformat()}")
        try:
            # The long-form response consumes most of the rate budget. Leave a
            # short cooldown before the next article rather than issuing a burst.
            if index:
                time.sleep(65)
            before_title_calls = title_generator.stats["ai_calls"]
            title_result = ai_result_with_retry(
                title_generator, [item], before_title_calls, "AI title generation"
            )
            title = title_result.output_data[-1]
            item["title"] = title

            before_content_calls = content_generator.stats["ai_calls"]
            content_result = ai_result_with_retry(
                content_generator, item, before_content_calls, "AI content generation"
            )
            content = content_result.output_data[0]

            image_plan = library.plan_article_images(item, title, content)
            inline_images = []
            for image in image_plan.inline_images:
                with_path = image.copy()
                with_path["path"] = str(library.library_dir / image["file"])
                with_path["marker"] = library.marker_for(image["file"])
                inline_images.append(with_path)

            publish_result = publisher.execute([{
                "title": title,
                "content": image_plan.content,
                "keyword": item["keyword"],
                "categories": [item["category"]],
                "publish_status": "future",
                "schedule_date": slot.isoformat(),
                "featured_image_path": str(library.library_dir / image_plan.featured_image["file"]) if image_plan.featured_image else None,
                "editorial_images": inline_images,
            }])
            if publish_result.status.value != "completed":
                raise RuntimeError(publish_result.error_message or "WordPress scheduling failed")
            post = publish_result.output_data[0]
            if not post.post_id or post.status != "future":
                raise RuntimeError(post.error or "WordPress did not return a future post")

            article_file = RESULTS_DIR / f"{index + 1:02d}_{post.post_id}.html"
            article_file.write_text(f"<!-- Title: {title} -->\n<!-- Keyword: {item['keyword']} -->\n{image_plan.content}", encoding="utf-8")
            results["articles"].append({
                "keyword": item["keyword"], "title": title, "post_id": post.post_id,
                "status": post.status, "url": post.url, "schedule_utc": slot.isoformat(),
                "featured_image": image_plan.featured_image.get("file") if image_plan.featured_image else None,
                "inline_images": [image["file"] for image in image_plan.inline_images],
                "local_article": str(article_file.relative_to(PROJECT_ROOT)),
            })
            save_results(results)
            log(f"Scheduled WordPress post #{post.post_id}: {title}")
        except Exception as exc:
            results["articles"].append({"keyword": item["keyword"], "status": "failed", "error": str(exc)})
            save_results(results)
            log(f"FAILED {item['keyword']}: {exc}")

    log("Scheduling run finished.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
