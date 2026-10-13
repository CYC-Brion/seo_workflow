"""
SEO工作流系统命令行接口
"""

import argparse
import os
import sys
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from config.settings import BASE_DIR
from src.utils import get_logger, config_manager
from src.modules import (
    ExcelProcessor,
    KeywordExpander,
    AIKeywordFilter,
    AITitleGenerator,
    AIContentGenerator,
    WordPressPublisher,
)
from src.core import WorkflowEngine, StateManager, WorkflowState, WorkflowStatus
from src.enhanced_features.image_integration import EditorialImageLibrary


logger = get_logger("cli")


def resolve_schedule_settings(
    schedule_hours: Optional[float],
    schedule_interval_hours: Optional[float],
    schedule_posts_enabled: bool,
    default_schedule_hours: float,
    default_schedule_interval_hours: float,
) -> Tuple[float, float]:
    """Resolve effective scheduling values for the workflow run."""
    effective_schedule_hours = (
        default_schedule_hours
        if schedule_hours is None and schedule_posts_enabled
        else (schedule_hours or 0)
    )
    effective_schedule_interval_hours = (
        default_schedule_interval_hours
        if schedule_interval_hours is None and effective_schedule_hours > 0
        else (schedule_interval_hours or 0)
    )
    return effective_schedule_hours, effective_schedule_interval_hours


def build_publish_schedule(
    total_posts: int,
    schedule_hours: float,
    schedule_interval_hours: float,
    now: Optional[datetime] = None,
) -> Tuple[str, List[Optional[str]]]:
    """Build the publish status and per-post schedule timestamps."""
    if schedule_hours <= 0:
        return "draft", [None] * total_posts

    base_time = (now or datetime.now(timezone.utc)) + timedelta(hours=schedule_hours)
    schedule_dates = [
        (base_time + timedelta(hours=schedule_interval_hours * index)).isoformat()
        for index in range(total_posts)
    ]
    return "future", schedule_dates


def _normalize_plan_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_keyword_metadata(file_path: Path) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
    """Load keyword sheet metadata and preserve workbook keyword order."""
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning(f"读取 keywords sheet 失败: {exc}")
        return {}, []

    if "keywords" not in workbook.sheetnames:
        return {}, []

    worksheet = workbook["keywords"]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return {}, []

    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    metadata_lookup: Dict[str, Dict[str, str]] = {}
    ordered_keywords: List[str] = []

    for raw_row in rows[1:]:
        row_data = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row_data[header] = raw_row[index] if index < len(raw_row) else None

        keyword = _normalize_plan_value(row_data.get("keyword") or row_data.get("keywords"))
        if not keyword:
            continue

        keyword_lower = keyword.lower()
        ordered_keywords.append(keyword)
        metadata_lookup[keyword_lower] = {
            "keyword": keyword,
            "category": _normalize_plan_value(row_data.get("category")),
            "lane": _normalize_plan_value(row_data.get("lane")),
            "priority": _normalize_plan_value(row_data.get("priority")),
            "cluster": _normalize_plan_value(row_data.get("cluster")),
            "internal_link_targets": _normalize_plan_value(row_data.get("internal_link_targets")),
            "notes": _normalize_plan_value(row_data.get("notes")),
        }

    return metadata_lookup, ordered_keywords


def load_workflow_plan(file_path: Path) -> List[Dict[str, str]]:
    """Load optional plan sheet for mixed lane execution."""
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning(f"读取 plan sheet 失败: {exc}")
        return []

    if "plan" not in workbook.sheetnames:
        return []

    worksheet = workbook["plan"]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    plan_rows: List[Dict[str, str]] = []

    for raw_row in rows[1:]:
        row_data = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row_data[header] = raw_row[index] if index < len(raw_row) else None

        keyword = _normalize_plan_value(row_data.get("keyword") or row_data.get("keywords"))
        if not keyword:
            continue

        direction = _normalize_plan_value(
            row_data.get("direction")
            or row_data.get("title_angle")
            or row_data.get("pillar_direction")
        )
        title_angle = _normalize_plan_value(
            row_data.get("title_angle")
            or row_data.get("direction")
            or row_data.get("pillar_direction")
        )
        plan_rows.append({
            "week": _normalize_plan_value(row_data.get("week")),
            "lane": _normalize_plan_value(row_data.get("lane")),
            "keyword": keyword,
            "category": _normalize_plan_value(row_data.get("category")),
            "priority": _normalize_plan_value(row_data.get("priority")),
            "cluster": _normalize_plan_value(row_data.get("cluster")),
            "direction": direction,
            "title_angle": title_angle,
            "internal_link_targets": _normalize_plan_value(row_data.get("internal_link_targets")),
        })

    return plan_rows


def build_default_title(keyword: str, direction: str = "") -> str:
    if direction:
        return f"Complete Guide to {keyword}: {direction}"
    return f"Complete Guide to {keyword}"


def build_work_items(
    filtered_keywords: List[str],
    keyword_categories: Dict[str, str],
    keyword_metadata: Optional[Dict[str, Dict[str, str]]] = None,
    ordered_keywords: Optional[List[str]] = None,
    plan_rows: Optional[List[Dict[str, str]]] = None,
) -> List[Dict[str, str]]:
    """Create the ordered workflow input items."""
    filtered_order = [keyword.strip() for keyword in filtered_keywords if keyword.strip()]
    filtered_lookup = {keyword.lower() for keyword in filtered_order}
    metadata_lookup = keyword_metadata or {}
    items: List[Dict[str, str]] = []
    planned_keywords = set()
    priority_order = {"a": 0, "b": 1, "c": 2}

    def build_item(
        keyword: str,
        category: str = "",
        lane: str = "",
        priority: str = "",
        cluster: str = "",
        direction: str = "",
        title_angle: str = "",
        internal_link_targets: str = "",
    ) -> Dict[str, str]:
        keyword_lower = keyword.lower()
        metadata = metadata_lookup.get(keyword_lower, {})
        from config.wordpress_config import DEFAULT_AI_CATEGORY

        resolved_category = (
            category.strip()
            or keyword_categories.get(keyword, "")
            or metadata.get("category", "")
            or DEFAULT_AI_CATEGORY
        )
        resolved_direction = direction.strip()
        resolved_title_angle = title_angle.strip() or resolved_direction

        return {
            "keyword": keyword,
            "title": build_default_title(keyword, resolved_direction),
            "category": resolved_category,
            "lane": lane.strip() or metadata.get("lane", "") or "ordinary",
            "priority": priority.strip() or metadata.get("priority", ""),
            "cluster": cluster.strip() or metadata.get("cluster", ""),
            "direction": resolved_direction,
            "title_angle": resolved_title_angle,
            "internal_link_targets": internal_link_targets.strip() or metadata.get("internal_link_targets", ""),
        }

    if plan_rows:
        for row in plan_rows:
            keyword = row.get("keyword", "").strip()
            if not keyword or keyword.lower() not in filtered_lookup:
                continue
            planned_keywords.add(keyword.lower())
            items.append(
                build_item(
                    keyword=keyword,
                    category=row.get("category", ""),
                    lane=row.get("lane", ""),
                    priority=row.get("priority", ""),
                    cluster=row.get("cluster", ""),
                    direction=row.get("direction", ""),
                    title_angle=row.get("title_angle", ""),
                    internal_link_targets=row.get("internal_link_targets", ""),
                )
            )

    ordinary_candidates: List[Tuple[int, int, Dict[str, str]]] = []
    keyword_position = {keyword.lower(): index for index, keyword in enumerate(filtered_order)}
    workbook_order = ordered_keywords or []
    seen_ordinary = set()

    for index, keyword in enumerate(workbook_order):
        keyword_lower = keyword.lower()
        if keyword_lower not in filtered_lookup or keyword_lower in planned_keywords:
            continue
        metadata = metadata_lookup.get(keyword_lower, {})
        lane = metadata.get("lane", "").strip().lower()
        if lane == "pillar":
            continue

        item = build_item(
            keyword=metadata.get("keyword", keyword),
            category=metadata.get("category", ""),
            lane=metadata.get("lane", "ordinary"),
            priority=metadata.get("priority", ""),
            cluster=metadata.get("cluster", ""),
            internal_link_targets=metadata.get("internal_link_targets", ""),
        )
        ordinary_candidates.append(
            (priority_order.get(item["priority"].lower(), 99), index, item)
        )
        seen_ordinary.add(keyword_lower)

    for keyword in filtered_order:
        keyword_lower = keyword.lower()
        if keyword_lower in planned_keywords or keyword_lower in seen_ordinary:
            continue

        item = build_item(keyword=keyword, lane="ordinary")
        ordinary_candidates.append(
            (
                priority_order.get(item["priority"].lower(), 99),
                len(workbook_order) + keyword_position.get(keyword_lower, 0),
                item,
            )
        )

    ordinary_candidates.sort(key=lambda candidate: (candidate[0], candidate[1]))
    items.extend(item for _, _, item in ordinary_candidates)

    return items


def setup_parser() -> argparse.ArgumentParser:
    """设置命令行参数解析器"""
    parser = argparse.ArgumentParser(
        description="SEO内容自动化工作流系统",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  %(prog)s config check                    # 检查配置
  %(prog)s excel read keywords.xlsx        # 读取Excel关键词
  %(prog)s workflow run keywords.xlsx      # 运行完整工作流
  %(prog)s --help                          # 显示帮助信息
        """
    )

    # 全局参数
    parser.add_argument(
        "-v", "--verbose",
        action="store_true",
        help="详细输出模式"
    )
    parser.add_argument(
        "--log-level",
        choices=["DEBUG", "INFO", "WARNING", "ERROR", "CRITICAL"],
        default="INFO",
        help="日志级别"
    )

    # 子命令
    subparsers = parser.add_subparsers(
        dest="command",
        title="可用命令",
        description="运行 '%(prog)s <command> --help' 获取命令详情"
    )

    # config 命令
    config_parser = subparsers.add_parser(
        "config",
        help="配置管理"
    )
    config_subparsers = config_parser.add_subparsers(dest="config_command")

    # config check
    config_check_parser = config_subparsers.add_parser(
        "check",
        help="检查配置状态"
    )
    config_check_parser.add_argument(
        "--test-api",
        action="store_true",
        help="测试API连接"
    )

    # config list
    config_list_parser = config_subparsers.add_parser(
        "list",
        help="列出配置信息"
    )
    config_list_parser.add_argument(
        "--format",
        choices=["json", "yaml", "table"],
        default="table",
        help="输出格式"
    )

    # excel 命令
    excel_parser = subparsers.add_parser(
        "excel",
        help="Excel文件处理"
    )
    excel_subparsers = excel_parser.add_subparsers(dest="excel_command")

    # excel read
    excel_read_parser = excel_subparsers.add_parser(
        "read",
        help="读取Excel文件中的关键词"
    )
    excel_read_parser.add_argument(
        "file",
        help="Excel文件路径"
    )
    excel_read_parser.add_argument(
        "-s", "--sheet",
        default=0,
        help="工作表名称或索引（默认: 0）"
    )
    excel_read_parser.add_argument(
        "-c", "--column",
        default="keywords",
        help="关键词列名（默认: 'keywords'）"
    )
    excel_read_parser.add_argument(
        "-o", "--output",
        help="输出文件路径（默认: 输出到控制台）"
    )
    excel_read_parser.add_argument(
        "--format",
        choices=["json", "csv", "txt"],
        default="txt",
        help="输出格式（默认: txt）"
    )

    # excel validate
    excel_validate_parser = excel_subparsers.add_parser(
        "validate",
        help="验证Excel文件"
    )
    excel_validate_parser.add_argument(
        "file",
        help="Excel文件路径"
    )

    # filter 命令
    filter_parser = subparsers.add_parser(
        "filter",
        help="AI关键词筛选"
    )
    filter_subparsers = filter_parser.add_subparsers(dest="filter_command")

    # filter run
    filter_run_parser = filter_subparsers.add_parser(
        "run",
        help="运行AI关键词筛选"
    )
    filter_run_parser.add_argument(
        "input",
        help="输入文件路径或关键词列表（JSON文件）"
    )
    filter_run_parser.add_argument(
        "-o", "--output",
        help="输出文件路径"
    )
    filter_run_parser.add_argument(
        "-b", "--batch-size",
        type=int,
        default=50,
        help="批量大小（默认: 50）"
    )
    filter_run_parser.add_argument(
        "-t", "--threshold",
        type=float,
        default=0.7,
        help="置信度阈值（默认: 0.7）"
    )
    filter_run_parser.add_argument(
        "--dry-run",
        action="store_true",
        help="试运行，不调用AI"
    )

    # filter stats
    filter_stats_parser = filter_subparsers.add_parser(
        "stats",
        help="显示筛选统计"
    )

    # workflow 命令
    workflow_parser = subparsers.add_parser(
        "workflow",
        help="工作流管理"
    )
    workflow_subparsers = workflow_parser.add_subparsers(dest="workflow_command")

    # workflow run
    workflow_run_parser = workflow_subparsers.add_parser(
        "run",
        help="运行工作流"
    )
    workflow_run_parser.add_argument(
        "input",
        help="输入文件路径"
    )
    workflow_run_parser.add_argument(
        "-o", "--output-dir",
        # 锚定项目根下的 output/，避免随启动目录漂移。注意这里刻意不用
        # config.settings.OUTPUT_DIR —— 那个指向 data/output，放的是中间研究产物；
        # scripts/schedule_pending_articles.py 与 build_article_keyword_queue.py
        # 都从 <root>/output 读 workflow_*.json，改动这里会让它们读不到。
        default=str(BASE_DIR / "output"),
        help=f"输出目录（默认: {BASE_DIR / 'output'}）"
    )
    workflow_run_parser.add_argument(
        "--dry-run",
        action="store_true",
        help="试运行，不调用AI和外部API"
    )
    workflow_run_parser.add_argument(
        "--use-preselected-keywords",
        action="store_true",
        help="输入文件已完成关键词研究与筛选：跳过扩词和AI关键词筛选，仅为表内关键词生成内容"
    )
    workflow_run_parser.add_argument(
        "--resume",
        help="从保存的状态恢复执行"
    )
    workflow_run_parser.add_argument(
        "--schedule-hours",
        type=float,
        default=None,
        help="定时发布：N小时后自动发布（0=立即存为草稿）。例如 --schedule-hours 24 表示24小时后发布"
    )

    workflow_run_parser.add_argument(
        "--schedule-interval-hours",
        type=float,
        default=None,
        help="定时发布时，每篇文章之间间隔 N 小时。0 表示同一时间发布"
    )

    # workflow status
    workflow_status_parser = workflow_subparsers.add_parser(
        "status",
        help="查看工作流状态"
    )
    workflow_status_parser.add_argument(
        "workflow_id",
        nargs="?",
        help="工作流ID（如果未提供，显示所有工作流）"
    )

    # workflow list
    workflow_list_parser = workflow_subparsers.add_parser(
        "list",
        help="列出工作流"
    )

    # workflow cancel
    workflow_cancel_parser = workflow_subparsers.add_parser(
        "cancel",
        help="取消工作流"
    )
    workflow_cancel_parser.add_argument(
        "workflow_id",
        help="工作流ID"
    )

    # workflow auto-publish
    workflow_auto_publish_parser = workflow_subparsers.add_parser(
        "auto-publish",
        help="自动发布草稿：持续运行，每间隔N分钟发1篇"
    )
    workflow_auto_publish_parser.add_argument(
        "--interval-minutes",
        type=int,
        default=60,
        help="发布间隔（分钟），默认60"
    )
    workflow_auto_publish_parser.add_argument(
        "--max-posts",
        type=int,
        default=0,
        help="最多发布篇数（0=不限）"
    )

    # workflow publish-next
    workflow_publish_next_parser = workflow_subparsers.add_parser(
        "publish-next",
        help="发布下一篇草稿（配合定时任务使用，每次运行发布1篇）"
    )

    # 工具命令
    tools_parser = subparsers.add_parser(
        "tools",
        help="工具函数"
    )
    tools_subparsers = tools_parser.add_subparsers(dest="tools_command")

    # tools version
    tools_version_parser = tools_subparsers.add_parser(
        "version",
        help="显示版本信息"
    )

    return parser


def handle_config(args: argparse.Namespace) -> int:
    """处理config命令"""
    if args.config_command == "check":
        return handle_config_check(args)
    elif args.config_command == "list":
        return handle_config_list(args)
    else:
        print("未知的config子命令")
        return 1


def handle_config_check(args: argparse.Namespace) -> int:
    """处理config check命令"""
    try:
        print("检查配置状态...")

        # 加载配置
        if not config_manager.load_all_configs():
            print("[FAIL] 配置加载失败")
            return 1

        print("[OK] 配置加载成功")

        # 检查API密钥
        api_status = config_manager.check_api_keys()
        print("\nAPI密钥状态:")
        for service, has_key in api_status.items():
            status = "[OK]" if has_key else "[FAIL]"
            print(f"  {service}: {status}")

        # 检查目录
        project_config = config_manager.project_config
        print("\n目录状态:")
        directories = [
            ("数据目录", project_config.data_dir),
            ("输入目录", project_config.input_dir),
            ("输出目录", project_config.output_dir),
            ("缓存目录", project_config.cache_dir),
            ("日志目录", project_config.logs_dir),
        ]

        for name, path in directories:
            if path.exists():
                print(f"  {name}: [OK] ({path})")
            else:
                print(f"  {name}: [FAIL] 不存在 ({path})")

        # 测试API连接（如果启用）
        if args.test_api:
            print("\n测试API连接...")
            # 这里可以添加实际的API测试
            print("  [WARN] API测试功能尚未实现")

        print("\n配置检查完成")
        return 0

    except Exception as e:
        print(f"[FAIL] 配置检查失败: {e}")
        logger.error("配置检查失败", exc_info=True)
        return 1


def handle_config_list(args: argparse.Namespace) -> int:
    """处理config list命令"""
    try:
        # 加载配置
        config_manager.load_all_configs()

        if args.format == "json":
            import json
            config_data = {
                "project": config_manager.project_config.__dict__,
                "ai_models": {
                    node: config.__dict__
                    for node, config in config_manager.ai_model_configs.items()
                },
                "wordpress": config_manager.wordpress_config.__dict__,
            }
            print(json.dumps(config_data, indent=2, default=str))
        else:
            # 表格格式
            print("项目配置:")
            print(f"  默认AI模型: {config_manager.project_config.default_ai_model}")
            print(f"  最大关键词批量: {config_manager.project_config.max_keywords_per_batch}")
            print(f"  文章字数: {config_manager.project_config.article_word_count}")

            print("\nAI模型配置:")
            for node_name, config in config_manager.ai_model_configs.items():
                print(f"  {node_name}: {config.provider} - {config.model}")

            print("\nWordPress配置:")
            wp_config = config_manager.wordpress_config
            print(f"  站点URL: {wp_config.url}")
            print(f"  用户名: {wp_config.username}")
            print(f"  默认状态: {wp_config.default_post_status}")

        return 0

    except Exception as e:
        print(f"[FAIL] 列出配置失败: {e}")
        logger.error("列出配置失败", exc_info=True)
        return 1


def handle_excel(args: argparse.Namespace) -> int:
    """处理excel命令"""
    if args.excel_command == "read":
        return handle_excel_read(args)
    elif args.excel_command == "validate":
        return handle_excel_validate(args)
    else:
        print("未知的excel子命令")
        return 1


def handle_excel_read(args: argparse.Namespace) -> int:
    """处理excel read命令"""
    try:
        file_path = Path(args.file)
        if not file_path.exists():
            print(f"[FAIL] 文件不存在: {file_path}")
            return 1

        print(f"读取Excel文件: {file_path}")

        # 创建Excel处理器
        processor = ExcelProcessor(
            node_id="excel_cli",
            config={
                "sheet_name": args.sheet,
                "keyword_column": args.column,
            }
        )

        # 读取关键词
        keywords = processor.execute_impl(file_path)
        stats = processor.get_statistics()

        print(f"[OK] 读取完成，共 {len(keywords)} 个关键词")
        print(f"  总行数: {stats['total_rows']}")
        print(f"  有效关键词: {stats['valid_keywords']}")
        print(f"  无效关键词: {stats['invalid_keywords']}")
        print(f"  重复关键词: {stats['duplicate_keywords']}")

        # 输出结果
        if args.output:
            output_path = Path(args.output)
            output_path.parent.mkdir(parents=True, exist_ok=True)

            if args.format == "json":
                with open(output_path, 'w', encoding='utf-8') as f:
                    json.dump(keywords, f, ensure_ascii=False, indent=2)
            elif args.format == "csv":
                import csv
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.writer(f)
                    writer.writerow(["keyword"])
                    for kw in keywords:
                        writer.writerow([kw])
            else:  # txt
                with open(output_path, 'w', encoding='utf-8') as f:
                    for kw in keywords:
                        f.write(f"{kw}\n")

            print(f"结果已保存到: {output_path}")
        else:
            # 输出到控制台
            print("\n关键词列表:")
            for i, kw in enumerate(keywords[:20], 1):  # 只显示前20个
                print(f"  {i}. {kw}")
            if len(keywords) > 20:
                print(f"  ... 还有 {len(keywords) - 20} 个关键词")

        return 0

    except Exception as e:
        print(f"[FAIL] 读取Excel失败: {e}")
        logger.error("读取Excel失败", exc_info=True)
        return 1


def handle_excel_validate(args: argparse.Namespace) -> int:
    """处理excel validate命令"""
    try:
        file_path = Path(args.file)
        if not file_path.exists():
            print(f"[FAIL] 文件不存在: {file_path}")
            return 1

        print(f"验证Excel文件: {file_path}")

        # 创建Excel处理器
        processor = ExcelProcessor(
            node_id="excel_validate_cli",
            config={}
        )

        # 验证文件
        is_valid = processor.validate_input(file_path)

        if is_valid:
            print("[OK] 文件验证通过")
            return 0
        else:
            print("[FAIL] 文件验证失败")
            return 1

    except Exception as e:
        print(f"[FAIL] 验证Excel失败: {e}")
        logger.error("验证Excel失败", exc_info=True)
        return 1


def handle_filter(args: argparse.Namespace) -> int:
    """处理filter命令"""
    if args.filter_command == "run":
        return handle_filter_run(args)
    elif args.filter_command == "stats":
        return handle_filter_stats(args)
    else:
        print("未知的filter子命令")
        return 1


def handle_filter_run(args: argparse.Namespace) -> int:
    """处理filter run命令"""
    try:
        input_path = Path(args.input)
        keywords = []

        # 加载输入数据
        if input_path.exists() and input_path.suffix.lower() == '.json':
            with open(input_path, 'r', encoding='utf-8') as f:
                keywords = json.load(f)
        elif input_path.exists():
            # 尝试作为文本文件读取
            with open(input_path, 'r', encoding='utf-8') as f:
                content = f.read().strip()
                if content.startswith('[') and content.endswith(']'):
                    keywords = json.loads(content)
                else:
                    keywords = [line.strip() for line in content.split('\n') if line.strip()]
        else:
            # 直接作为关键词字符串
            if args.input.startswith('[') and args.input.endswith(']'):
                keywords = json.loads(args.input)
            else:
                keywords = [args.input]

        print(f"开始AI关键词筛选，共 {len(keywords)} 个关键词")
        print(f"批量大小: {args.batch_size}, 置信度阈值: {args.threshold}")

        if args.dry_run:
            print("试运行模式，跳过AI调用")
            print(f"将处理 {len(keywords)} 个关键词")
            return 0

        # 创建筛选器
        filter_node = AIKeywordFilter(
            node_id="filter_cli",
            config={
                "batch_size": args.batch_size,
                "confidence_threshold": args.threshold,
            }
        )

        # 执行筛选
        filtered_keywords = filter_node.execute_impl(keywords)
        stats = filter_node.get_statistics()

        print(f"[OK] 筛选完成")
        print(f"  原始关键词: {stats['total_keywords']}")
        print(f"  相关关键词: {stats['relevant_keywords']}")
        print(f"  不相关关键词: {stats['irrelevant_keywords']}")
        print(f"  低置信度关键词: {stats['low_confidence_keywords']}")
        print(f"  AI调用次数: {stats['ai_calls']}")
        print(f"  总Tokens: {stats['total_tokens']}")
        print(f"  总成本: ${stats['total_cost']:.4f}")
        print(f"  最终关键词: {len(filtered_keywords)}")

        # 输出结果
        if args.output:
            output_path = Path(args.output)
            output_path.parent.mkdir(parents=True, exist_ok=True)

            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(filtered_keywords, f, ensure_ascii=False, indent=2)

            print(f"结果已保存到: {output_path}")
        else:
            # 输出到控制台
            print("\n筛选后的关键词:")
            for i, kw in enumerate(filtered_keywords[:20], 1):
                print(f"  {i}. {kw}")
            if len(filtered_keywords) > 20:
                print(f"  ... 还有 {len(filtered_keywords) - 20} 个关键词")

        return 0

    except Exception as e:
        print(f"[FAIL] AI关键词筛选失败: {e}")
        logger.error("AI关键词筛选失败", exc_info=True)
        return 1


def handle_filter_stats(args: argparse.Namespace) -> int:
    """处理filter stats命令"""
    print("筛选统计功能尚未实现")
    return 0


def handle_workflow(args: argparse.Namespace) -> int:
    """处理workflow命令"""
    if args.workflow_command == "run":
        return handle_workflow_run(args)
    elif args.workflow_command == "status":
        return handle_workflow_status(args)
    elif args.workflow_command == "list":
        return handle_workflow_list(args)
    elif args.workflow_command == "cancel":
        return handle_workflow_cancel(args)
    elif args.workflow_command == "auto-publish":
        return handle_workflow_auto_publish(args)
    elif args.workflow_command == "publish-next":
        return handle_workflow_publish_next(args)
    else:
        print("未知的workflow子命令")
        return 1


def _save_workflow_state(engine: WorkflowEngine, state_manager: StateManager, name: str) -> Optional[str]:
    """保存工作流状态到磁盘，返回保存的 workflow_id"""
    try:
        result = engine.get_result()
        if result:
            state = state_manager.create_state_from_result(result)
            state.workflow_name = name
            state_manager.save_workflow_state(state)
            return result.workflow_id
    except Exception as e:
        logger.warning(f"保存工作流状态失败: {e}")
    return None


def handle_workflow_run(args: argparse.Namespace) -> int:
    """处理workflow run命令 — 执行完整的6节点SEO内容自动化管线"""
    try:
        input_path = Path(args.input)
        if not input_path.exists():
            print(f"[FAIL] 输入文件不存在: {input_path}")
            return 1

        output_dir = Path(args.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        dry_run = args.dry_run
        use_preselected_keywords = getattr(args, "use_preselected_keywords", False)
        if dry_run:
            print("=" * 60)
            print("  DRY-RUN 模式 — 跳过 AI 调用和 WordPress 发布")
            print("=" * 60)

        # 初始化 StateManager
        state_manager = StateManager()

        # 创建 WorkflowEngine（用于状态跟踪）
        engine = WorkflowEngine("seo-workflow-cli")
        print(f"\n工作流引擎已创建: {engine.workflow_id}\n")

        # ============================
        # 步骤 1: Excel 读取关键词
        # ============================
        step = 1
        total_steps = 6
        print(f"[{step}/{total_steps}] 读取 Excel 关键词...")
        # Operational workbooks place an executive Overview before the machine-
        # readable queue. Prefer the conventional ``keywords`` sheet when it is
        # available, while retaining first-sheet compatibility for ad-hoc files.
        input_sheet_name: Any = 0
        try:
            input_workbook = load_workbook(input_path, read_only=True, data_only=True)
            if "keywords" in input_workbook.sheetnames:
                input_sheet_name = "keywords"
            input_workbook.close()
        except Exception:
            pass
        excel_processor = ExcelProcessor("excel_reader", config={"sheet_name": input_sheet_name})
        result = excel_processor.execute(input_path)
        if result.status.value != "completed":
            print(f"  [FAIL] Excel 读取失败: {result.error_message}")
            return 1
        keywords: list = result.output_data
        # 提取关键词→分类映射
        original_categories: Dict[str, str] = getattr(excel_processor, 'category_map', {})
        if original_categories:
            cats = set(original_categories.values())
            print(f"  [OK] 读取到 {len(keywords)} 个关键词, 分类: {cats}")
        else:
            print(f"  [OK] 读取到 {len(keywords)} 个关键词")
        if keywords:
            print(f"  预览: {keywords[:5]}{'...' if len(keywords) > 5 else ''}")

        keyword_metadata, ordered_keywords = load_keyword_metadata(input_path)
        if keyword_metadata:
            print(f"  [OK] keywords sheet metadata: {len(keyword_metadata)} 条")

        plan_rows = load_workflow_plan(input_path)
        if plan_rows:
            print(f"  [OK] plan sheet: {len(plan_rows)} 条混合排期项")

        # ============================
        # 步骤 2: 关键词扩展
        # ============================
        step += 1
        from config.wordpress_config import DEFAULT_AI_CATEGORY
        if use_preselected_keywords:
            expanded_keywords = keywords
            keyword_categories = {
                keyword: original_categories.get(keyword, DEFAULT_AI_CATEGORY)
                for keyword in expanded_keywords
            }
            print(f"\n[{step}/{total_steps}] 扩展关键词...")
            print("  [SKIP] 使用已精选词库：不再调用关键词扩展服务")
        else:
            print(f"\n[{step}/{total_steps}] 扩展关键词...")
            expander_config = {
                "max_expansions_per_keyword": 10,
                "request_delay_seconds": 2,
                "enable_google_browser_fallback": os.getenv("ENABLE_GOOGLE_BROWSER_FALLBACK", "false").strip().lower() in {"1", "true", "yes", "on"},
                "browser_fallback_timeout_seconds": float(os.getenv("GOOGLE_BROWSER_FALLBACK_TIMEOUT_SECONDS", "20")),
            }
            http_proxy = os.getenv("HTTP_PROXY")
            if http_proxy:
                expander_config["http_proxy"] = http_proxy
            expander = KeywordExpander("keyword_expander", config=expander_config)
            result = expander.execute(keywords)
            if result.status.value != "completed":
                print(f"  [FAIL] 关键词扩展失败: {result.error_message}")
                return 1
            expanded_keywords = result.output_data

            # 扩展关键词继承父关键词的分类
            keyword_categories = {}
            for ekw in expanded_keywords:
                assigned = False
                for okw, cat in original_categories.items():
                    if ekw == okw or ekw.startswith(okw + " "):
                        keyword_categories[ekw] = cat
                        assigned = True
                        break
                if not assigned:
                    keyword_categories[ekw] = DEFAULT_AI_CATEGORY

            expander_stats = expander.get_statistics()
            print(f"  [OK] 扩展完成: {len(keywords)} → {len(expanded_keywords)} 个关键词")
            print(f"  Google API: {expander_stats.get('api_requests', 0)}, "
                  f"Google Suggest: {expander_stats.get('google_suggest_requests', 0)}, "
                  f"Bing Suggest: {expander_stats.get('bing_suggest_requests', 0)}")
            planner_ok = expander_stats.get("google_ads_keyword_planner_requests", 0)
            planner_failed = expander_stats.get("google_ads_keyword_planner_failures", 0)
            print(f"  Google Ads Keyword Planner: {planner_ok} 成功, {planner_failed} 失败")
            if planner_failed and not planner_ok:
                print("  [WARN]  Google Ads 关键词规划师全部失败——本次扩展只剩 Bing Suggest 供货，")
                print("          关键词会明显偏向本地化/低商业意图（如 'xxx near me'）。")
                print("          常见原因：本地代理未启动（.env 的 HTTP_PROXY），或 Google Ads developer token 仍是 Test Access。")

        # ============================
        # 步骤 3: AI 关键词筛选
        # ============================
        step += 1
        print(f"\n[{step}/{total_steps}] AI 关键词筛选...")
        if use_preselected_keywords:
            print("  [SKIP] 使用已精选词库：不再调用AI关键词筛选")
            filtered_keywords = expanded_keywords
        elif dry_run:
            print("  [SKIP]  DRY-RUN: 跳过AI筛选，保留所有关键词")
            filtered_keywords = expanded_keywords
        else:
            filter_node = AIKeywordFilter("keyword_filter", config={
                "batch_size": 50,
                "confidence_threshold": 0.7,
                "enable_fallback": True,
            })
            result = filter_node.execute(expanded_keywords)
            if result.status.value != "completed":
                print(f"  [WARN]  AI筛选未完全成功: {result.error_message}")
                filtered_keywords = expanded_keywords
            else:
                filtered_keywords = result.output_data
            filter_stats = filter_node.get_statistics()
            print(f"  相关: {filter_stats.get('relevant_keywords', '?')}, "
                  f"不相关: {filter_stats.get('irrelevant_keywords', '?')}")
        print(f"  [OK] 筛选后: {len(filtered_keywords)} 个关键词")

        if not filtered_keywords:
            print("  [FAIL] 筛选后无关键词，工作流终止")
            return 1

        workflow_items = build_work_items(
            filtered_keywords,
            keyword_categories,
            keyword_metadata=keyword_metadata,
            ordered_keywords=ordered_keywords,
            plan_rows=plan_rows,
        )
        print(f"  [OK] 生成 {len(workflow_items)} 个最终内容任务")

        # ============================
        # 步骤 4: AI 标题生成
        # ============================
        step += 1
        print(f"\n[{step}/{total_steps}] AI 标题生成...")
        if dry_run:
            print("  [SKIP]  DRY-RUN: 使用默认标题")
            titles = [build_default_title(item["keyword"], item.get("direction", "")) for item in workflow_items]
        else:
            title_generator = AITitleGenerator("title_generator", config={
                "batch_size": 10,
                "max_titles_per_keyword": 1,
                "enable_fallback": True,
            })
            result = title_generator.execute(workflow_items)
            if result.status.value != "completed":
                print(f"  [WARN]  标题生成未完全成功，使用默认标题")
                titles = [build_default_title(item["keyword"], item.get("direction", "")) for item in workflow_items]
            else:
                titles = result.output_data
        print(f"  [OK] 生成 {len(titles)} 个标题")
        if titles:
            print(f"  预览: {titles[0][:80]}...")

        # ============================
        # 步骤 5: AI 内容生成
        # ============================
        step += 1
        print(f"\n[{step}/{total_steps}] AI 内容生成...")
        # 构建 content_generator 的输入格式
        content_inputs = []
        for i, item in enumerate(workflow_items):
            title = titles[i] if i < len(titles) else item.get("title", build_default_title(item["keyword"], item.get("direction", "")))
            cat = item.get("category", "") or keyword_categories.get(item["keyword"], DEFAULT_AI_CATEGORY)
            content_inputs.append({
                "keyword": item["keyword"],
                "title": title,
                "category": cat,
                "direction": item.get("direction", ""),
                "lane": item.get("lane", ""),
                "priority": item.get("priority", ""),
                "cluster": item.get("cluster", ""),
                "title_angle": item.get("title_angle", ""),
                "internal_link_targets": item.get("internal_link_targets", ""),
            })

        if dry_run:
            print("  [SKIP]  DRY-RUN: 使用默认内容")
            content_generator = AIContentGenerator("content_generator", config={
                "batch_size": 3,
                "target_word_count": 1500,
                "html_format": True,
                "enable_fallback": True,
            })
            contents = [
                content_generator._generate_default_content(
                    inp["keyword"],
                    inp["title"],
                    direction=inp.get("direction", ""),
                    lane=inp.get("lane", ""),
                )
                for inp in content_inputs
            ]
        else:
            content_generator = AIContentGenerator("content_generator", config={
                "batch_size": 3,
                "target_word_count": 1500,
                "html_format": True,
                "enable_fallback": True,
            })
            result = content_generator.execute(content_inputs)
            if result.status.value != "completed":
                print(f"  [WARN]  内容生成未完全成功，使用默认内容")
                contents = [
                    f"<article><h2>{inp['title']}</h2><p>Fallback content for {inp['keyword']}.</p></article>"
                    for inp in content_inputs
                ]
            else:
                contents = result.output_data
        print(f"  [OK] 生成 {len(contents)} 篇文章")
        for i, c in enumerate(contents[:2]):
            print(f"  文章 {i+1}: {len(c)} 字符, 预览: {c[:100]}...")

        article_meta = getattr(content_generator, '_article_meta', {}) if not dry_run else {}
        article_metas = [article_meta.get(content, {}) for content in contents]

        # Select images only after the final article structure is available, so body
        # evidence can be placed beside the section it actually supports.
        image_library = EditorialImageLibrary()
        article_image_plans = []
        for index, content in enumerate(contents):
            item = workflow_items[index] if index < len(workflow_items) else {}
            title = titles[index] if index < len(titles) else item.get("title", "")
            image_plan = image_library.plan_article_images(item, title, content)
            contents[index] = image_plan.content
            article_image_plans.append(image_plan)
        featured_count = sum(
            1 for image_plan in article_image_plans if image_plan.featured_image
        )
        inline_count = sum(
            len(image_plan.inline_images) for image_plan in article_image_plans
        )
        print(
            f"  [OK] 图片策划: {featured_count} 张首图, "
            f"{inline_count} 张正文证据图"
        )

        # ============================
        # 步骤 6: WordPress 发布
        # ============================
        step += 1
        print(f"\n[{step}/{total_steps}] WordPress 发布...")

        wp_runtime_config = config_manager.wordpress_config
        schedule_hours, schedule_interval_hours = resolve_schedule_settings(
            schedule_hours=getattr(args, "schedule_hours", None),
            schedule_interval_hours=getattr(args, "schedule_interval_hours", None),
            schedule_posts_enabled=wp_runtime_config.schedule_posts,
            default_schedule_hours=wp_runtime_config.default_schedule_hours,
            default_schedule_interval_hours=wp_runtime_config.default_schedule_interval_hours,
        )
        publish_status, post_schedule_dates = build_publish_schedule(
            total_posts=len(titles),
            schedule_hours=schedule_hours,
            schedule_interval_hours=schedule_interval_hours,
        )
        if post_schedule_dates and post_schedule_dates[0]:
            print(f"  [SCHEDULE] first post will publish at {post_schedule_dates[0][:19]} UTC")
            if schedule_interval_hours > 0 and len(post_schedule_dates) > 1:
                print(
                    f"  [SCHEDULE] {len(post_schedule_dates)} posts queued, "
                    f"last post at {post_schedule_dates[-1][:19]} UTC, "
                    f"interval {schedule_interval_hours} hour(s)"
                )

        wp_inputs = []
        for i, title in enumerate(titles):
            workflow_item = workflow_items[i] if i < len(workflow_items) else {}
            kw = workflow_item.get("keyword", "")
            content = contents[i] if i < len(contents) else ""
            meta = article_metas[i] if i < len(article_metas) else {}
            cat_name = workflow_item.get("category", "") or keyword_categories.get(kw, DEFAULT_AI_CATEGORY)
            post_schedule_date = post_schedule_dates[i] if i < len(post_schedule_dates) else None
            image_plan = article_image_plans[i] if i < len(article_image_plans) else None
            featured_image = image_plan.featured_image if image_plan else None
            inline_images = []
            if image_plan:
                for image in image_plan.inline_images:
                    image_with_marker = image.copy()
                    image_with_marker["marker"] = image_library.marker_for(image["file"])
                    # 引用 WP 媒体库已有条目。没有引用才回填本地路径走上传兜底
                    # —— 但图库已经只选出有引用的资产，所以这条兜底实际是死路，
                    # 这正是目的：不再产生重复副本。
                    reference = image_library.wordpress_reference(image)
                    if reference:
                        image_with_marker["media_id"] = reference["id"]
                        image_with_marker["source_url"] = reference["source_url"]
                    else:
                        image_with_marker["path"] = str(
                            image_library.library_dir / image["file"]
                        )
                    inline_images.append(image_with_marker)

            featured_reference = image_library.wordpress_reference(featured_image or {})
            featured_image_path = None
            if featured_image and not featured_reference:
                featured_image_path = str(image_library.library_dir / featured_image["file"])

            wp_inputs.append({
                "title": title,
                "content": content,
                "keyword": kw,
                "slug": meta.get("slug", ""),
                "excerpt": meta.get("excerpt", ""),
                "categories": [cat_name],
                "publish_status": publish_status,
                "schedule_date": post_schedule_date,
                "featured_media_id": (
                    featured_reference["id"] if featured_reference else None
                ),
                "featured_image_path": featured_image_path,
                "editorial_images": inline_images,
            })

        if dry_run:
            print("  [SKIP]  DRY-RUN: 跳过 WordPress 发布")
            print(f"  将发布 {len(wp_inputs)} 篇文章（{publish_status} 状态）")
            publish_results = [
                {"post_id": None, "title": inp["title"], "status": publish_status, "url": "(dry-run)"}
                for inp in wp_inputs
            ]
        else:
            try:
                publisher = WordPressPublisher("wordpress_publisher", config={
                    "batch_size": 5,
                    "api_type": "rest",
                    "default_status": "draft",
                    "auto_create_categories": False,
                    "auto_create_tags": False,
                    "enable_media_upload": True,
                })
                result = publisher.execute(wp_inputs)
                if result.status.value == "completed":
                    # 将 dataclass 转为 dict 以便后续统计
                    publish_results = [
                        {"post_id": r.post_id, "title": r.title, "status": r.status, "url": r.url}
                        for r in result.output_data
                    ]
                else:
                    print(f"  [WARN]  发布未完全成功: {result.error_message}")
                    publish_results = [
                        {"post_id": None, "title": inp["title"], "status": "failed", "url": ""}
                        for inp in wp_inputs
                    ]
            except Exception as wp_err:
                print(f"  [WARN]  WordPress 连接失败: {wp_err}")
                print(f"  文章内容已生成，可稍后手动发布")
                publish_results = [
                    {"post_id": None, "title": inp["title"], "status": "skipped", "url": ""}
                    for inp in wp_inputs
                ]

        # 统计发布结果
        published = sum(1 for r in publish_results if r.get("status") in ("publish", "published"))
        scheduled = sum(1 for r in publish_results if r.get("status") == "future")
        drafts = sum(1 for r in publish_results if r.get("status") == "draft")
        failed = sum(1 for r in publish_results if r.get("status") in ("failed", "skipped"))
        print(f"  [OK] 发布完成: {published} 已发布, {scheduled} 已排期, {drafts} 草稿, {failed} 跳过/失败")

        # ============================
        # 保存结果
        # ============================
        print(f"\n保存结果到 {output_dir}...")
        output_data = {
            "workflow_id": engine.workflow_id,
            "input_file": str(input_path),
            "dry_run": dry_run,
            "steps": {
                "1_excel_keywords": {"count": len(keywords), "data": keywords},
                "2_expanded_keywords": {"count": len(expanded_keywords), "data": expanded_keywords},
                "3_filtered_keywords": {"count": len(filtered_keywords), "data": filtered_keywords},
                "4_titles": {"count": len(titles), "data": titles},
                "5_contents": {"count": len(contents),
                               "data": [c[:200] + "..." for c in contents]},  # 截断预览
                "6_publish_results": publish_results,
            },
        }
        results_file = output_dir / f"workflow_{engine.workflow_id}.json"
        with open(results_file, "w", encoding="utf-8") as f:
            json.dump(output_data, f, ensure_ascii=False, indent=2)
        print(f"  [OK] 结果已保存: {results_file}")

        # 保存完整文章到单独文件
        articles_dir = output_dir / "articles"
        articles_dir.mkdir(exist_ok=True)
        for i, (inp, content) in enumerate(zip(wp_inputs, contents)):
            safe_title = "".join(c for c in inp["title"][:50] if c.isalnum() or c in " _-").strip()
            article_file = articles_dir / f"{i+1:03d}_{safe_title}.html"
            with open(article_file, "w", encoding="utf-8") as f:
                f.write(f"<!-- Title: {inp['title']} -->\n")
                f.write(f"<!-- Keyword: {inp['keyword']} -->\n")
                f.write(content)
        print(f"  [OK] {len(contents)} 篇文章已保存: {articles_dir}")

        # 保存工作流状态
        state_manager = StateManager()
        workflow_id = _save_workflow_state(engine, state_manager, "seo-workflow-cli")

        # ============================
        # 打印摘要
        # ============================
        print(f"\n{'=' * 60}")
        print(f"  工作流执行完成!")
        print(f"  Workflow ID: {engine.workflow_id}")
        print(f"  关键词: {len(keywords)} → 扩展 {len(expanded_keywords)} → 筛选 {len(filtered_keywords)}")
        print(f"  生成: {len(titles)} 标题, {len(contents)} 篇文章")
        print(f"  发布: {published} 已发布, {scheduled} 已排期, {drafts} 草稿, {failed} 跳过")
        print(f"  输出目录: {output_dir}")
        print(f"{'=' * 60}")

        return 0

    except Exception as e:
        print(f"[FAIL] 工作流运行失败: {e}")
        logger.error("工作流运行失败", exc_info=True)
        return 1


def handle_workflow_status(args: argparse.Namespace) -> int:
    """处理workflow status命令 — 查看最近一次工作流状态"""
    try:
        state_manager = StateManager()

        if args.workflow_id:
            state = state_manager.load_workflow_state(args.workflow_id)
            if not state:
                print(f"[FAIL] 未找到工作流: {args.workflow_id}")
                return 1
            states_to_show = [state]
        else:
            all_states = state_manager.list_workflow_states(limit=1)
            if not all_states:
                print("没有找到已保存的工作流状态")
                print("运行 'python -m src.cli.main workflow run <excel文件>' 来执行工作流")
                return 0
            # 取最新的一条
            latest = all_states[0]
            state = state_manager.load_workflow_state(latest.get("workflow_id", ""))
            if not state:
                print("[FAIL] 无法加载工作流状态")
                return 1
            states_to_show = [state]

        for s in states_to_show:
            print(f"\n工作流状态")
            print(f"  ID:       {s.workflow_id}")
            print(f"  名称:     {s.workflow_name}")
            print(f"  状态:     {s.status.value}")
            print(f"  开始时间: {s.start_time}")
            print(f"  结束时间: {s.end_time}")
            print(f"  耗时:     {s.execution_time:.1f}s" if s.execution_time else "")
            if s.node_states:
                print(f"  节点状态:")
                for node_id, node_info in s.node_states.items():
                    status = node_info.get("status", "unknown")
                    icon = {"completed": "[OK]", "failed": "[FAIL]", "running": "[RUN]"}.get(status, "[WAIT]")
                    print(f"    {icon} {node_id}: {status}")

        return 0

    except Exception as e:
        print(f"[FAIL] 查询工作流状态失败: {e}")
        logger.error("查询工作流状态失败", exc_info=True)
        return 1


def handle_workflow_list(args: argparse.Namespace) -> int:
    """处理workflow list命令 — 列出所有已保存的工作流"""
    try:
        state_manager = StateManager()
        states = state_manager.list_workflow_states(limit=50)

        if not states:
            print("没有找到已保存的工作流状态")
            return 0

        print(f"\n{'ID':<38} {'名称':<25} {'状态':<12} {'时间'}")
        print("-" * 95)
        for s in states:
            wid = s.get("workflow_id", "")[:36]
            name = s.get("workflow_name", "")[:23]
            status = s.get("status", "")
            start = s.get("start_time", "")
            if isinstance(start, str) and len(start) > 19:
                start = start[:19]
            print(f"{wid:<38} {name:<25} {status:<12} {start}")

        print(f"\n共 {len(states)} 条记录")
        return 0

    except Exception as e:
        print(f"[FAIL] 列出工作流失败: {e}")
        logger.error("列出工作流失败", exc_info=True)
        return 1


def handle_workflow_cancel(args: argparse.Namespace) -> int:
    """处理workflow cancel命令 — 取消指定工作流"""
    try:
        state_manager = StateManager()
        state = state_manager.load_workflow_state(args.workflow_id)

        if not state:
            print(f"[FAIL] 未找到工作流: {args.workflow_id}")
            return 1

        if state.status == WorkflowStatus.COMPLETED:
            print(f"工作流 {args.workflow_id} 已经完成，无法取消")
            return 0
        elif state.status == WorkflowStatus.CANCELLED:
            print(f"工作流 {args.workflow_id} 已经被取消")
            return 0

        state.status = WorkflowStatus.CANCELLED
        state_manager.save_workflow_state(state)
        print(f"[OK] 工作流 {args.workflow_id} 已取消")
        return 0

    except Exception as e:
        print(f"[FAIL] 取消工作流失败: {e}")
        logger.error("取消工作流失败", exc_info=True)
        return 1


def handle_workflow_auto_publish(args: argparse.Namespace) -> int:
    """处理workflow auto-publish命令 — 自动定时发布草稿"""
    import time as time_mod
    from datetime import datetime, timezone
    from src.modules.wordpress_integration.wp_api_client import get_wordpress_client

    interval = max(args.interval_minutes, 1)
    max_posts = args.max_posts
    published_count = 0

    print("=" * 60)
    print("  草稿自动发布器")
    print(f"  间隔: {interval} 分钟")
    if max_posts > 0:
        print(f"  上限: {max_posts} 篇")
    print("=" * 60)

    try:
        client = get_wordpress_client(api_type="rest")
    except Exception as e:
        print(f"[FAIL] WordPress 连接失败: {e}")
        return 1

    try:
        while True:
            # 获取当前草稿
            try:
                drafts = client.get_posts(status="draft", per_page=100)
            except Exception as e:
                print(f"[WARN] 获取草稿列表失败: {e}，60秒后重试...")
                time_mod.sleep(60)
                continue

            if not drafts:
                print(f"\n[OK] 没有草稿了！共发布 {published_count} 篇")
                break

            # 已经达到上限
            if max_posts > 0 and published_count >= max_posts:
                print(f"\n[OK] 已达到上限 {max_posts} 篇，停止。剩余草稿: {len(drafts)}")
                break

            # 取最旧的草稿发布
            draft = drafts[0]
            post_id = draft.get("id")
            title = draft.get("title", {}).get("rendered", "Untitled")

            print(f"\n[{datetime.now(timezone.utc).strftime('%H:%M:%S')}] "
                  f"发布草稿 #{post_id}: {title[:60]}...")

            try:
                client.update_post(post_id, {"status": "publish"})
                published_count += 1
                print(f"  [OK] #{post_id} 已发布 ({published_count} 篇)")
            except Exception as e:
                print(f"  [FAIL] #{post_id} 发布失败: {e}，跳过")
                continue

            # 检查是否还有剩余
            if max_posts == 0 or published_count < max_posts:
                try:
                    remaining = client.get_posts(status="draft", per_page=1)
                except Exception:
                    remaining = drafts[1:]  # 粗略估计
                if not remaining:
                    print(f"\n[OK] 全部发布完毕！共 {published_count} 篇")
                    break

            # 已到上限则停止
            if max_posts > 0 and published_count >= max_posts:
                print(f"\n[OK] 已完成 {published_count} 篇，停止")
                break

            # 等待间隔
            next_publish = datetime.now(timezone.utc).timestamp() + interval * 60
            next_time = datetime.fromtimestamp(next_publish).strftime('%H:%M:%S')
            print(f"  下一篇: {next_time} (等待 {interval} 分钟...)")
            time_mod.sleep(interval * 60)

        return 0

    except KeyboardInterrupt:
        print(f"\n\n[STOP] 用户中断。已发布 {published_count} 篇")
        return 0
    except Exception as e:
        print(f"[FAIL] 自动发布异常: {e}")
        logger.error("自动发布异常", exc_info=True)
        return 1


def handle_workflow_publish_next(args: argparse.Namespace) -> int:
    """处理workflow publish-next命令 — 发布1篇草稿（配合定时任务）"""
    from datetime import datetime, timezone
    from src.modules.wordpress_integration.wp_api_client import get_wordpress_client

    try:
        client = get_wordpress_client(api_type="rest")
    except Exception as e:
        print(f"[FAIL] WordPress 连接失败: {e}")
        return 1

    try:
        drafts = client.get_posts(status="draft", per_page=1)
        if not drafts:
            print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M')}] 没有草稿，跳过")
            return 0

        draft = drafts[0]
        post_id = draft.get("id")
        title = draft.get("title", {}).get("rendered", "Untitled")

        client.update_post(post_id, {"status": "publish"})
        print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M')}] "
              f"已发布 #{post_id}: {title[:60]}")
        return 0

    except Exception as e:
        print(f"[FAIL] 发布失败: {e}")
        return 1


def handle_tools(args: argparse.Namespace) -> int:
    """处理tools命令"""
    if args.tools_command == "version":
        return handle_tools_version(args)
    else:
        print("未知的tools子命令")
        return 1


def handle_tools_version(args: argparse.Namespace) -> int:
    """处理tools version命令"""
    try:
        # 从pyproject.toml读取版本信息
        import toml
        pyproject_path = Path(__file__).parent.parent.parent / "pyproject.toml"

        if pyproject_path.exists():
            with open(pyproject_path, 'r', encoding='utf-8') as f:
                pyproject = toml.load(f)
                version = pyproject.get('project', {}).get('version', '未知')
                name = pyproject.get('project', {}).get('name', 'seo-content-automation')
                description = pyproject.get('project', {}).get('description', '')

                print(f"{name} v{version}")
                print(description)
        else:
            print("SEO内容自动化工作流系统 v0.1.0")

        return 0
    except Exception:
        print("SEO内容自动化工作流系统 v0.1.0")
        return 0


def main(argv: Optional[list] = None) -> int:
    """CLI主入口点"""
    parser = setup_parser()
    args = parser.parse_args(argv)

    # 如果没有提供命令，显示帮助
    if not args.command:
        parser.print_help()
        return 0

    # 设置日志级别
    logger.level = args.log_level

    try:
        # 根据命令分发处理
        if args.command == "config":
            return handle_config(args)
        elif args.command == "excel":
            return handle_excel(args)
        elif args.command == "filter":
            return handle_filter(args)
        elif args.command == "workflow":
            return handle_workflow(args)
        elif args.command == "tools":
            return handle_tools(args)
        else:
            print(f"未知命令: {args.command}")
            return 1

    except KeyboardInterrupt:
        print("\n操作被用户中断")
        return 130
    except Exception as e:
        if args.verbose:
            import traceback
            traceback.print_exc()
        else:
            print(f"[FAIL] 执行失败: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
